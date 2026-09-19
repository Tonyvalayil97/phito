import io
import os
import re

import fitz  # PyMuPDF
import pandas as pd
import pdfplumber
import pytesseract
import streamlit as st
from docx import Document
from openai import OpenAI
from PIL import Image

MODEL = "Qwen/Qwen3-4B-Instruct-2507"
MAX_CONTEXT_CHARS = 18000

st.set_page_config(page_title="AI Document Intelligence", page_icon="📊", layout="wide")

st.markdown("""
<style>
.block-container {max-width: 1200px; padding-top: 2rem;}
[data-testid="stFileUploader"] {border: 1px dashed #7c8cff; border-radius: 14px; padding: 1rem;}
.hero {padding: 1.5rem; border-radius: 18px; background: linear-gradient(135deg,#18223d,#4656a6); color:white; margin-bottom:1rem;}
.hero p {margin-bottom:0; opacity:.9}
</style>
<div class="hero">
  <h1>📊 AI Document Intelligence</h1>
  <p>Read documents, extract tables, create charts, and ask AI for grounded insights.</p>
</div>
""", unsafe_allow_html=True)


def get_token():
    try:
        return st.secrets["HF_TOKEN"]
    except (KeyError, FileNotFoundError):
        return os.getenv("HF_TOKEN", "")


def unique_headers(values):
    seen, headers = {}, []
    for index, value in enumerate(values):
        base = str(value).strip() if value not in (None, "") else f"Column {index + 1}"
        seen[base] = seen.get(base, 0) + 1
        headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return headers


def table_to_dataframe(raw_table):
    if not raw_table or len(raw_table) < 2:
        return None
    width = max(len(row) for row in raw_table if row)
    rows = [list(row) + [None] * (width - len(row)) for row in raw_table if row]
    if len(rows) < 2:
        return None
    frame = pd.DataFrame(rows[1:], columns=unique_headers(rows[0]))
    frame = frame.replace({None: ""}).dropna(how="all")
    return frame if not frame.empty else None


@st.cache_data(show_spinner=False)
def extract_pdf(data: bytes):
    doc = fitz.open(stream=data, filetype="pdf")
    pages, ocr_pages = [], 0
    for number, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        if len(text) < 30:
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.open(io.BytesIO(pix.tobytes("png")))
            text = pytesseract.image_to_string(image).strip()
            ocr_pages += 1
        pages.append(f"--- Page {number} ---\n{text}")

    tables = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            for table_number, raw_table in enumerate(page.extract_tables(), start=1):
                frame = table_to_dataframe(raw_table)
                if frame is not None:
                    tables.append({
                        "name": f"Page {page_number} · Table {table_number}",
                        "page": page_number,
                        "data": frame,
                    })
    return "\n\n".join(pages), tables, ocr_pages, len(doc)


@st.cache_data(show_spinner=False)
def extract_image(data: bytes):
    image = Image.open(io.BytesIO(data)).convert("RGB")
    return pytesseract.image_to_string(image).strip()


@st.cache_data(show_spinner=False)
def extract_docx(data: bytes):
    doc = Document(io.BytesIO(data))
    blocks = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = []
    for number, table in enumerate(doc.tables, start=1):
        raw = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        frame = table_to_dataframe(raw)
        if frame is not None:
            tables.append({"name": f"Word table {number}", "page": None, "data": frame})
            blocks.append(frame.to_csv(index=False))
    return "\n".join(blocks), tables


def extract(upload):
    data, name = upload.getvalue(), upload.name.lower()
    if name.endswith(".pdf"):
        text, tables, ocr_pages, total_pages = extract_pdf(data)
        return text, tables, f"{total_pages} page(s); OCR used on {ocr_pages}; {len(tables)} table(s)"
    if name.endswith((".png", ".jpg", ".jpeg", ".tiff", ".bmp")):
        return extract_image(data), [], "Image processed with OCR"
    if name.endswith(".docx"):
        text, tables = extract_docx(data)
        return text, tables, f"Word document parsed; {len(tables)} table(s)"
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="replace"), [], "Text file read"
    raise ValueError("Unsupported file type")


def clean_numeric(series):
    cleaned = series.astype(str).str.replace(r"[^0-9.\-()]", "", regex=True)
    cleaned = cleaned.str.replace(r"^\((.*)\)$", r"-\1", regex=True)
    return pd.to_numeric(cleaned, errors="coerce")


def prepare_chart_frame(frame):
    chart_frame = frame.copy()
    numeric_columns = []
    for column in chart_frame.columns:
        converted = clean_numeric(chart_frame[column])
        if converted.notna().sum() >= max(2, len(frame) // 2):
            chart_frame[column] = converted
            numeric_columns.append(column)
    return chart_frame, numeric_columns


def tables_as_context(tables):
    sections = []
    for item in tables:
        sections.append(f"TABLE: {item['name']}\n{item['data'].head(100).to_csv(index=False)}")
    return "\n\n".join(sections)


def ask_model(instruction, document_text, tables, token):
    client = OpenAI(base_url="https://router.huggingface.co/v1", api_key=token)
    combined = f"{document_text}\n\n{tables_as_context(tables)}"[:MAX_CONTEXT_CHARS]
    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": (
                "You are a careful document and data analyst. Use only the supplied document "
                "and tables. Mention relevant table names, values, comparisons and trends. "
                "If evidence is missing, say so clearly."
            )},
            {"role": "user", "content": f"DOCUMENT DATA:\n{combined}\n\nTASK:\n{instruction}"},
        ],
        temperature=0.2,
        max_tokens=900,
    )
    return response.choices[0].message.content


with st.sidebar:
    st.header("Settings")
    st.caption(f"Model: {MODEL}")
    token = get_token()
    st.success("HF token detected") if token else st.warning("Add HF_TOKEN in Streamlit secrets")
    st.markdown("Supported: **PDF, scanned PDF, images, DOCX, TXT**")
    st.caption("Structured tables require a digital PDF or DOCX. Scanned pages are still readable through OCR.")

upload = st.file_uploader(
    "Upload a document",
    type=["pdf", "png", "jpg", "jpeg", "tiff", "bmp", "docx", "txt"],
)

if not upload:
    st.info("Upload a PDF containing a table to try summaries, table extraction and charts.")
    st.stop()

try:
    with st.spinner("Reading text, running OCR and detecting tables..."):
        text, tables, details = extract(upload)
except Exception as exc:
    st.error(f"Could not read this document: {exc}")
    st.stop()

if not text.strip() and not tables:
    st.error("No readable text or tables were found. Try a clearer scan.")
    st.stop()

st.success(f"Analysis ready — {details}")
metric1, metric2, metric3, metric4 = st.columns(4)
metric1.metric("Characters", f"{len(text):,}")
metric2.metric("Words", f"{len(text.split()):,}")
metric3.metric("Tables", len(tables))
metric4.download_button(
    "Download text", text, file_name=f"{upload.name.rsplit('.', 1)[0]}_text.txt",
    mime="text/plain", use_container_width=True,
)

tab_summary, tab_tables, tab_charts, tab_text = st.tabs(
    ["AI Summary & Q&A", "Extracted Tables", "Charts", "Extracted Text"]
)

with tab_summary:
    if "messages" not in st.session_state:
        st.session_state.messages = []
    col1, col2, col3, col4 = st.columns(4)
    summarize = col1.button("Document summary", use_container_width=True)
    table_summary = col2.button("Table insights", use_container_width=True, disabled=not tables)
    risks = col3.button("Risks & actions", use_container_width=True)
    clear = col4.button("Clear chat", use_container_width=True)
    if clear:
        st.session_state.messages = []
        st.rerun()
    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])
    prompt = st.chat_input("Ask about the document, numbers, tables or trends")
    instruction = (
        "Summarize the document with clear headings, key numbers and conclusions."
        if summarize else
        "Analyze every extracted table. Explain major values, comparisons, trends, anomalies and conclusions."
        if table_summary else
        "Identify risks, anomalies, decisions and recommended action items supported by the document."
        if risks else prompt
    )
    if instruction:
        if not token:
            st.error("Add HF_TOKEN in Streamlit Cloud → App settings → Secrets.")
        else:
            st.session_state.messages.append({"role": "user", "content": instruction})
            with st.chat_message("user"):
                st.markdown(instruction)
            with st.chat_message("assistant"):
                with st.spinner("Analyzing the document and tables..."):
                    try:
                        answer = ask_model(instruction, text, tables, token)
                        st.markdown(answer)
                        st.session_state.messages.append({"role": "assistant", "content": answer})
                    except Exception as exc:
                        st.error(f"Model request failed: {exc}")

with tab_tables:
    if not tables:
        st.info("No structured tables were detected. Scanned tables appear in extracted text but may require a specialized table-OCR service for rows and columns.")
    else:
        selected_name = st.selectbox("Choose a table", [item["name"] for item in tables])
        selected = next(item for item in tables if item["name"] == selected_name)
        frame = selected["data"]
        st.dataframe(frame, use_container_width=True, hide_index=True)
        st.download_button(
            "Download this table as CSV", frame.to_csv(index=False).encode("utf-8"),
            file_name=re.sub(r"[^A-Za-z0-9_-]+", "_", selected_name) + ".csv",
            mime="text/csv",
        )
        numeric = frame.apply(clean_numeric)
        stats = numeric.dropna(axis=1, how="all").describe().T
        if not stats.empty:
            st.subheader("Automatic numeric statistics")
            st.dataframe(stats, use_container_width=True)

with tab_charts:
    if not tables:
        st.info("Charts become available when a structured table is detected.")
    else:
        chart_table_name = st.selectbox("Table", [item["name"] for item in tables], key="chart_table")
        chart_item = next(item for item in tables if item["name"] == chart_table_name)
        chart_frame, numeric_columns = prepare_chart_frame(chart_item["data"])
        if not numeric_columns:
            st.warning("This table has no reliably numeric column to graph.")
        else:
            category_options = list(chart_frame.columns)
            default_category = next((c for c in category_options if c not in numeric_columns), category_options[0])
            control1, control2, control3 = st.columns(3)
            x_column = control1.selectbox("Category / X-axis", category_options, index=category_options.index(default_category))
            y_column = control2.selectbox("Numeric value / Y-axis", numeric_columns)
            chart_type = control3.selectbox("Chart type", ["Bar", "Line", "Area"])
            plot = chart_frame[[x_column, y_column]].dropna(subset=[y_column]).head(100).set_index(x_column)
            if chart_type == "Bar":
                st.bar_chart(plot)
            elif chart_type == "Line":
                st.line_chart(plot)
            else:
                st.area_chart(plot)
            st.caption("Charts show up to the first 100 usable rows from the selected table.")

with tab_text:
    st.text_area("OCR / extracted text", text, height=550)
